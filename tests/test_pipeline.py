import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import docx
from infolumi.documents import read_references, write_references
from infolumi.pipeline import ReferencePipeline
from infolumi.pubmed_client import PubMedClient


class PipelineTests(unittest.IsolatedAsyncioTestCase):
    async def test_document_to_citations_preserves_missing_positions(self):
        chunks = []

        class Parser:
            async def parse(self, refs):
                chunks.append(refs)
                return (
                    json.dumps([{"title": ref} for ref in refs])
                    if len(refs) == 5
                    else "invalid"
                )

        searches = []

        def search(ref):
            searches.append(ref["title"])
            return ref["title"]

        def fetch(ids):
            return {i: {"ArticleTitle": "Title " + i} for i in reversed(ids)}

        pipeline = ReferencePipeline(
            Parser(), SimpleNamespace(search=search, fetch=fetch), search_delay=0
        )
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            document = docx.Document()
            for i in range(7):
                document.add_paragraph(str(i))
            document.save(source)
            progress = AsyncMock()
            results = await pipeline.process(source, progress)
            self.assertEqual([len(c) for c in chunks], [5, 2])
            self.assertEqual(searches, ["0", "1", "2", "3", "4"])
            self.assertIn("Title 0", results[0])
            self.assertEqual(results[-2:], ["6. Not Found", "7. Not Found"])
            self.assertEqual(progress.await_count, 2)
            for suffix in ["docx", "txt"]:
                target = Path(directory) / ("out." + suffix)
                write_references(results, target)
                actual = (
                    read_references(target)
                    if suffix == "docx"
                    else target.read_text().splitlines()
                )
                self.assertEqual(actual, results)

    async def test_empty_document(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.docx"
            docx.Document().save(path)
            parser = SimpleNamespace(parse=AsyncMock())
            self.assertEqual(
                await ReferencePipeline(
                    parser, SimpleNamespace(fetch=lambda ids: {})
                ).process(path),
                [],
            )
            parser.parse.assert_not_awaited()


class PubMedTests(unittest.TestCase):
    def test_search_and_fetch_without_network(self):
        calls = []

        def get(url, **kwargs):
            calls.append(kwargs)
            return SimpleNamespace(json=lambda: {"result": {"uids": [{"pubmed": "7"}]}})

        def post(url, **kwargs):
            return SimpleNamespace(
                raise_for_status=lambda: None,
                text="<PubmedArticleSet><PubmedArticle><MedlineCitation><PMID>7</PMID><Article><ArticleTitle>T</ArticleTitle></Article></MedlineCitation></PubmedArticle></PubmedArticleSet>",
            )

        client = PubMedClient(SimpleNamespace(get=get, post=post))
        self.assertEqual(client.search({"title": "T"}), "7")
        self.assertEqual(calls[0]["timeout"], 30)
        self.assertEqual(client.fetch(["7"])["7"]["ArticleTitle"], "T")
        self.assertEqual(client.fetch([]), {})


if __name__ == "__main__":
    unittest.main()
