"""Read reference paragraphs and write DOCX or text documents."""

from pathlib import Path

import docx


def read_references(path):
    return [paragraph.text for paragraph in docx.Document(path).paragraphs]


def write_references(references, path):
    path = Path(path)
    if path.suffix.lower() == ".docx":
        document = docx.Document()
        for reference in references:
            document.add_paragraph(reference)
        document.save(path)
    elif path.suffix.lower() == ".txt":
        path.write_text("\n".join(references), encoding="utf-8")
    else:
        raise ValueError("Output must be DOCX or TXT")
